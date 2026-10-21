import re
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def build_docx(md_path, output_paths):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = docx.Document()

    # Configuración de márgenes estándar (2.54 cm / 1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilos de fuente base
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(4)

    # --- PORTADA ---
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pre.paragraph_format.space_before = Pt(80)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t = p_title.add_run("SafeVoice: Plataforma Web para la Prevención, el Reporte Anónimo y la Gestión de Casos de Bullying y Ciberbullying en Instituciones Educativas\n")
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1B, 0x5E, 0x9E)

    run_sub = p_title.add_run("Documento Acumulativo de Avance 2")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x2C, 0x5F, 0x57)
    p_title.paragraph_format.space_after = Pt(120)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.line_spacing = 1.4

    runs_meta = [
        ("Mariana Arbeláez, Daniel López y Esteban Álvarez\n\n", True, 12, RGBColor(0x33, 0x33, 0x33)),
        ("Facultad de Ingeniería, Universidad Pontificia Bolivariana\n", False, 11, RGBColor(0x55, 0x55, 0x55)),
        ("Proyecto Aplicado en TIC (2026-2)\n", False, 11, RGBColor(0x55, 0x55, 0x55)),
        ("Profesora Yuri Marcela Escobar Cardona\n\n", False, 11, RGBColor(0x55, 0x55, 0x55)),
        ("18 de octubre de 2026", True, 11, RGBColor(0x1B, 0x5E, 0x9E)),
    ]
    for text, bold, sz, color in runs_meta:
        r = p_meta.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.color.rgb = color

    doc.add_page_break()

    # --- PARSING MARKDOWN LINES ---
    lines = md_text.split('\n')
    i = 0
    in_table = False
    table_lines = []
    in_code_block = False

    while i < len(lines):
        line = lines[i].strip()

        # Omitir metadata inicial ya plasmada en la portada
        if i < 15 and (line.startswith('# SafeVoice') or line.startswith('*') or line.startswith('**') or line.startswith('Facultad') or line.startswith('Medellín')):
            i += 1
            continue

        # Separadores horizontales
        if line == '---':
            i += 1
            continue

        # Bloques de código (Mermaid / JSON / Gherkin)
        if line.startswith('```'):
            if in_code_block:
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            # Renderizar línea de código como párrafo preformateado
            if line:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(line)
                r.font.name = 'Consolas'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # Detección de tablas markdown
        if line.startswith('|') and line.endswith('|'):
            table_lines.append(line)
            i += 1
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Procesar la tabla recolectada
            process_markdown_table(doc, table_lines)
            table_lines = []
            continue

        # Encabezados
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(clean_md_formatting(line[2:]))
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1B, 0x5E, 0x9E)
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(clean_md_formatting(line[3:]))
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x2C, 0x5F, 0x57)
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(clean_md_formatting(line[4:]))
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1B, 0x5E, 0x9E)
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(clean_md_formatting(line[5:]))
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, line[2:])
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            match = re.match(r'^\d+\.\s*(.*)', line)
            add_formatted_runs(p, match.group(1))
        elif line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(clean_md_formatting(line[2:]))
            r.font.italic = True
            r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        elif line:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_formatted_runs(p, line)

        i += 1

    # Guardar en las rutas requeridas
    for out_path in output_paths:
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        doc.save(out_path)
        print(f"Documento Word guardado exitosamente en: {out_path}")

def clean_md_formatting(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    return text.strip()

def add_formatted_runs(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('*') and part.endswith('*'):
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        elif part.startswith('`') and part.endswith('`'):
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)

def process_markdown_table(doc, raw_lines):
    rows = []
    for l in raw_lines:
        if re.match(r'^\s*\|?\s*:?-+:?\s*(\|?\s*:?-+:?\s*)+\|?\s*$', l):
            continue
        cells = [c.strip() for c in l.strip('|').split('|')]
        rows.append(cells)

    if not rows:
        return

    num_cols = max(len(r) for r in rows)
    num_rows = len(rows)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "D0D7DE")

    for r_idx, row_data in enumerate(rows):
        is_header = (r_idx == 0)
        table_row = table.rows[r_idx]
        
        # Repetir encabezado en cambio de página
        if is_header:
            trPr = table_row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c_idx in range(num_cols):
            cell = table_row.cells[c_idx]
            val = row_data[c_idx] if c_idx < len(row_data) else ""
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1

            if is_header:
                set_cell_background(cell, "1B5E9E")
                r = p.add_run(clean_md_formatting(val))
                r.bold = True
                r.font.name = 'Calibri'
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                bg_color = "F6F8FA" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg_color)
                # Formatear celdas con estado OK / Error / Pendiente
                clean_val = clean_md_formatting(val)
                p.text = ""
                r = p.add_run(clean_val)
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                if clean_val == "OK" or "✓" in clean_val:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x15, 0x80, 0x3D) # Verde
                elif clean_val == "Error" or clean_val == "✗":
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xBE, 0x12, 0x3C) # Rojo
                elif clean_val == "Pendiente":
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xB4, 0x53, 0x09) # Ámbar

    doc.add_paragraph() # Espacio tras tabla

if __name__ == '__main__':
    md_file = 'docs/SafeVoice_Avance2_Documento_Acumulativo.md'
    out_files = [
        os.path.join(os.environ.get('USERPROFILE', ''), 'Downloads', 'SafeVoice_Avance2_Documento_Acumulativo.docx'),
        'docs/SafeVoice_Avance2_Documento_Acumulativo.docx'
    ]
    build_docx(md_file, out_files)
